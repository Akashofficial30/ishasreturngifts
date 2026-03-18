from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages
from django.db.models import Sum, Count, Q
from django.utils import timezone
from django.http import HttpResponse
from datetime import timedelta
from products.models import Product, Category
from orders.models import Order, OrderItem
from payments.models import Payment
from orders.email_utils import send_order_confirmation_email, send_admin_order_notification
from django.contrib.auth.models import User
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def is_admin(user):
    return user.is_authenticated and user.is_staff


def get_new_msg_count():
    try:
        from contact.models import ContactMessage
        return ContactMessage.objects.filter(status='new').count()
    except Exception:
        return 0


# ── AUTH ─────────────────────────────────────────────────
def dashboard_login(request):
    if request.user.is_authenticated and request.user.is_staff:
        return redirect('dashboard_home')
    if request.method == 'POST':
        user = authenticate(request, username=request.POST.get('username'), password=request.POST.get('password'))
        if user and user.is_staff:
            login(request, user)
            return redirect('dashboard_home')
        messages.error(request, 'Invalid credentials or not an admin.')
    return render(request, 'dashboard/login.html')


def dashboard_logout(request):
    logout(request)
    return redirect('dashboard_login')


# ── HOME ─────────────────────────────────────────────────
@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def dashboard_home(request):
    today = timezone.now().date()
    last_30 = timezone.now() - timedelta(days=30)

    total_orders = Order.objects.count()
    online_revenue = Order.objects.filter(payment_status='paid', payment_method='online').aggregate(t=Sum('total_price'))['t'] or 0
    cod_revenue = Order.objects.filter(payment_method='cod', order_status='delivered').aggregate(t=Sum('total_price'))['t'] or 0
    total_revenue = float(online_revenue) + float(cod_revenue)
    total_products = Product.objects.filter(is_active=True).count()
    total_customers = User.objects.filter(is_staff=False).count()
    pending_orders = Order.objects.filter(order_status='pending').count()
    today_orders = Order.objects.filter(created_at__date=today).count()
    today_revenue = Order.objects.filter(created_at__date=today, payment_status='paid').aggregate(t=Sum('total_price'))['t'] or 0
    monthly_revenue = Order.objects.filter(created_at__gte=last_30, payment_status='paid').aggregate(t=Sum('total_price'))['t'] or 0
    recent_orders = Order.objects.prefetch_related('items').order_by('-created_at')[:8]
    low_stock = Product.objects.filter(stock_quantity__lte=10, is_active=True).order_by('stock_quantity')[:5]
    order_status_data = list(Order.objects.values('order_status').annotate(count=Count('id')))
    revenue_chart = []
    for i in range(6, -1, -1):
        day = timezone.now() - timedelta(days=i)
        rev = Order.objects.filter(created_at__date=day.date(), payment_status='paid').aggregate(t=Sum('total_price'))['t'] or 0
        revenue_chart.append({'day': day.strftime('%a'), 'revenue': float(rev)})
    top_products = OrderItem.objects.values('product_name').annotate(total_qty=Sum('quantity'), total_revenue=Sum('product_price')).order_by('-total_qty')[:5]

    context = {
        'total_orders': total_orders,
        'total_revenue': total_revenue,
        'total_products': total_products,
        'total_customers': total_customers,
        'pending_orders': pending_orders,
        'today_orders': today_orders,
        'today_revenue': today_revenue,
        'monthly_revenue': monthly_revenue,
        'cod_revenue': cod_revenue,
        'online_revenue': online_revenue,
        'recent_orders': recent_orders,
        'low_stock': low_stock,
        'order_status_data': order_status_data,
        'revenue_chart': revenue_chart,
        'top_products': top_products,
        'pending_count': pending_orders,
        'new_msg_count': get_new_msg_count(),
    }
    return render(request, 'dashboard/home.html', context)


# ── ORDERS ───────────────────────────────────────────────
@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def order_list(request):
    orders = Order.objects.prefetch_related('items').order_by('-created_at')
    status = request.GET.get('status', '')
    payment = request.GET.get('payment', '')
    search = request.GET.get('search', '')
    method = request.GET.get('method', '')
    if status: orders = orders.filter(order_status=status)
    if payment: orders = orders.filter(payment_status=payment)
    if method: orders = orders.filter(payment_method=method)
    if search:
        orders = orders.filter(Q(customer_name__icontains=search) | Q(customer_phone__icontains=search) | Q(payment_id__icontains=search))
    return render(request, 'dashboard/orders.html', {
        'orders': orders, 'status_filter': status, 'payment_filter': payment,
        'method_filter': method, 'search': search, 'total_count': orders.count(),
        'status_choices': Order.STATUS_CHOICES, 'payment_choices': Order.PAYMENT_STATUS_CHOICES,
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def order_detail(request, order_id):
    order = get_object_or_404(Order, id=order_id)
    if request.method == 'POST':
        action = request.POST.get('action', 'update')
        if action == 'resend_mail':
            send_order_confirmation_email(order)
            send_admin_order_notification(order)
            messages.success(request, f'Emails resent for Order #{order.get_short_order_id()}!')
        else:
            order.order_status = request.POST.get('order_status', order.order_status)
            order.payment_status = request.POST.get('payment_status', order.payment_status)
            order.notes = request.POST.get('notes', '')
            order.save()
            messages.success(request, f'Order #{order.get_short_order_id()} updated!')
        return redirect('dashboard_order_detail', order_id=order.id)
    return render(request, 'dashboard/order_detail.html', {
        'order': order,
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


# ── EXCEL EXPORTS ────────────────────────────────────────
@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def export_orders_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="6B1F2A")
    thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    headers = ['Order ID','Customer','Phone','Email','Address','City','State','Pincode','Items','Total (₹)','Payment Method','Payment Status','Order Status','Date']
    widths = [12,20,14,28,35,14,14,10,40,12,16,14,14,20]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center'); cell.border = thin
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 25
    gold_fill = PatternFill("solid", fgColor="FFF8E1")
    for row_num, order in enumerate(Order.objects.prefetch_related('items').order_by('-created_at'), 2):
        items_str = '; '.join([f"{i.product_name} x{i.quantity}" for i in order.items.all()])
        row_data = [f"#{order.get_short_order_id()}", order.customer_name, order.customer_phone,
                    order.customer_email or '—', order.address, order.city, order.state, order.pincode,
                    items_str, float(order.total_price), order.get_payment_method_display(),
                    order.payment_status.title(), order.order_status.title(),
                    order.created_at.strftime('%d %b %Y %H:%M')]
        fill = gold_fill if row_num % 2 == 0 else None
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.border = thin; cell.alignment = Alignment(vertical='center', wrap_text=True)
            if fill: cell.fill = fill
        ws.row_dimensions[row_num].height = 18
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="ishas_orders.xlsx"'
    wb.save(response)
    return response


@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def export_products_excel(request):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Products"
    header_font = Font(bold=True, color="FFFFFF"); header_fill = PatternFill("solid", fgColor="6B1F2A")
    thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    headers = ['Name','Category','Price (₹)','Offer Price (₹)','Stock','Featured','Active','Created']
    widths = [30,20,12,14,8,10,8,20]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center'); cell.border = thin
        ws.column_dimensions[get_column_letter(i)].width = w
    for row_num, p in enumerate(Product.objects.select_related('category').all(), 2):
        row = [p.name, p.category.name, float(p.price), float(p.offer_price) if p.offer_price else '—',
               p.stock_quantity, 'Yes' if p.is_featured else 'No', 'Yes' if p.is_active else 'No',
               p.created_at.strftime('%d %b %Y')]
        for col, val in enumerate(row, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.border = thin; cell.alignment = Alignment(vertical='center')
        ws.row_dimensions[row_num].height = 16
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="ishas_products.xlsx"'
    wb.save(response); return response


@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def export_customers_excel(request):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Customers"
    header_font = Font(bold=True, color="FFFFFF"); header_fill = PatternFill("solid", fgColor="6B1F2A")
    thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    headers = ['Username','Full Name','Email','Orders','Total Spent (₹)','Joined']
    widths = [18,22,30,8,16,20]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center'); cell.border = thin
        ws.column_dimensions[get_column_letter(i)].width = w
    for row_num, u in enumerate(User.objects.filter(is_staff=False), 2):
        orders = Order.objects.filter(customer_email=u.email)
        spent = orders.filter(payment_status='paid').aggregate(t=Sum('total_price'))['t'] or 0
        row = [u.username, u.get_full_name() or '—', u.email, orders.count(), float(spent), u.date_joined.strftime('%d %b %Y')]
        for col, val in enumerate(row, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.border = thin; cell.alignment = Alignment(vertical='center')
        ws.row_dimensions[row_num].height = 16
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="ishas_customers.xlsx"'
    wb.save(response); return response


# ── PRODUCTS ─────────────────────────────────────────────
@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def product_list(request):
    products = Product.objects.select_related('category').order_by('-created_at')
    search = request.GET.get('search', '')
    category = request.GET.get('category', '')
    stock = request.GET.get('stock', '')
    if search: products = products.filter(Q(name__icontains=search))
    if category: products = products.filter(category__slug=category)
    if stock == 'low': products = products.filter(stock_quantity__lte=10)
    elif stock == 'out': products = products.filter(stock_quantity=0)
    return render(request, 'dashboard/products.html', {
        'products': products, 'categories': Category.objects.all(),
        'search': search, 'category_filter': category, 'stock_filter': stock,
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def product_add(request):
    categories = Category.objects.all()
    if request.method == 'POST':
        from django.utils.text import slugify
        name = request.POST.get('name')
        slug = slugify(name)
        base_slug, counter = slug, 1
        while Product.objects.filter(slug=slug).exists():
            slug = f"{base_slug}-{counter}"; counter += 1
        product = Product(name=name, slug=slug, category_id=request.POST.get('category'),
                          description=request.POST.get('description'), price=request.POST.get('price'),
                          offer_price=request.POST.get('offer_price') or None,
                          stock_quantity=request.POST.get('stock_quantity', 0),
                          is_featured=request.POST.get('is_featured') == 'on',
                          is_active=request.POST.get('is_active') == 'on')
        if request.FILES.get('image'): product.image = request.FILES['image']
        product.save()
        messages.success(request, f'Product "{name}" added!')
        return redirect('dashboard_products')
    return render(request, 'dashboard/product_form.html', {
        'categories': categories, 'action': 'Add',
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def product_edit(request, product_id):
    product = get_object_or_404(Product, id=product_id)
    categories = Category.objects.all()
    if request.method == 'POST':
        product.name = request.POST.get('name')
        product.category_id = request.POST.get('category')
        product.description = request.POST.get('description')
        product.price = request.POST.get('price')
        product.offer_price = request.POST.get('offer_price') or None
        product.stock_quantity = request.POST.get('stock_quantity', 0)
        product.is_featured = request.POST.get('is_featured') == 'on'
        product.is_active = request.POST.get('is_active') == 'on'
        if request.FILES.get('image'): product.image = request.FILES['image']
        product.save()
        messages.success(request, f'Product "{product.name}" updated!')
        return redirect('dashboard_products')
    return render(request, 'dashboard/product_form.html', {
        'product': product, 'categories': categories, 'action': 'Edit',
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def product_delete(request, product_id):
    product = get_object_or_404(Product, id=product_id)
    if request.method == 'POST':
        name = product.name; product.delete()
        messages.success(request, f'"{name}" deleted.')
        return redirect('dashboard_products')
    return render(request, 'dashboard/confirm_delete.html', {
        'object': product, 'type': 'Product',
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


# ── CATEGORIES ───────────────────────────────────────────
@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def category_list(request):
    categories = Category.objects.annotate(product_count=Count('products'))
    return render(request, 'dashboard/categories.html', {
        'categories': categories,
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def category_add(request):
    if request.method == 'POST':
        from django.utils.text import slugify
        name = request.POST.get('name')
        slug = slugify(name)
        base_slug, counter = slug, 1
        while Category.objects.filter(slug=slug).exists():
            slug = f"{base_slug}-{counter}"; counter += 1
        cat = Category(name=name, slug=slug, description=request.POST.get('description', ''))
        if request.FILES.get('image'): cat.image = request.FILES['image']
        cat.save()
        messages.success(request, f'Category "{name}" created!')
        return redirect('dashboard_categories')
    return render(request, 'dashboard/category_form.html', {
        'action': 'Add',
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def category_edit(request, cat_id):
    category = get_object_or_404(Category, id=cat_id)
    if request.method == 'POST':
        category.name = request.POST.get('name')
        category.description = request.POST.get('description', '')
        if request.FILES.get('image'): category.image = request.FILES['image']
        category.save()
        messages.success(request, 'Category updated!')
        return redirect('dashboard_categories')
    return render(request, 'dashboard/category_form.html', {
        'category': category, 'action': 'Edit',
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def category_delete(request, cat_id):
    category = get_object_or_404(Category, id=cat_id)
    if request.method == 'POST':
        category.delete()
        messages.success(request, 'Category deleted.')
        return redirect('dashboard_categories')
    return render(request, 'dashboard/confirm_delete.html', {
        'object': category, 'type': 'Category',
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


# ── CUSTOMERS ────────────────────────────────────────────
@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def customer_list(request):
    customers = User.objects.filter(is_staff=False).order_by('-date_joined')
    search = request.GET.get('search', '')
    if search:
        customers = customers.filter(Q(username__icontains=search) | Q(email__icontains=search))
    customer_data = []
    for c in customers:
        orders = Order.objects.filter(customer_email=c.email)
        customer_data.append({
            'user': c,
            'order_count': orders.count(),
            'total_spent': orders.filter(payment_status='paid').aggregate(t=Sum('total_price'))['t'] or 0,
            'last_order': orders.order_by('-created_at').first(),
        })
    return render(request, 'dashboard/customers.html', {
        'customer_data': customer_data, 'search': search,
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


# ── PAYMENTS ─────────────────────────────────────────────
@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def payment_list(request):
    payments = Payment.objects.select_related('order').order_by('-created_at')
    status = request.GET.get('status', '')
    if status: payments = payments.filter(status=status)
    total_collected = Payment.objects.filter(status='paid').aggregate(t=Sum('amount'))['t'] or 0
    cod_collected = Order.objects.filter(payment_method='cod', order_status='delivered').aggregate(t=Sum('total_price'))['t'] or 0
    return render(request, 'dashboard/payments.html', {
        'payments': payments, 'status_filter': status,
        'total_collected': total_collected, 'cod_collected': cod_collected,
        'total_combined': float(total_collected) + float(cod_collected),
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


# ── CONTACT MESSAGES ────────────────────────────────────
@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def contact_messages(request):
    from contact.models import ContactMessage
    msgs = ContactMessage.objects.all()
    status = request.GET.get('status', '')
    subject = request.GET.get('subject', '')
    if status: msgs = msgs.filter(status=status)
    if subject: msgs = msgs.filter(subject=subject)

    if request.method == 'POST':
        msg_id = request.POST.get('msg_id')
        new_status = request.POST.get('status')
        notes = request.POST.get('admin_notes', '')
        if msg_id:
            try:
                cm = ContactMessage.objects.get(id=msg_id)
                if new_status: cm.status = new_status
                cm.admin_notes = notes
                cm.save()
                messages.success(request, 'Message updated!')
            except ContactMessage.DoesNotExist:
                pass
        return redirect('dashboard_contact_messages')

    return render(request, 'dashboard/contact_messages.html', {
        'contact_messages': msgs,
        'status_filter': status,
        'subject_filter': subject,
        'new_count': ContactMessage.objects.filter(status='new').count(),
        'pending_count': Order.objects.filter(order_status='pending').count(),
        'new_msg_count': get_new_msg_count(),
    })


# ── IMPORT EXCEL ─────────────────────────────────────────
@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def import_products_excel(request):
    from .import_utils import import_products_from_excel
    if request.method == 'POST':
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            messages.error(request, 'Please select an Excel file to upload.')
            return redirect('dashboard_products')
        if not excel_file.name.endswith(('.xlsx', '.xls')):
            messages.error(request, 'Only .xlsx or .xls files are supported.')
            return redirect('dashboard_products')
        success, errors = import_products_from_excel(excel_file)
        if success:
            messages.success(request, f'✅ Successfully imported {success} product{"s" if success > 1 else ""}!')
        if errors:
            for err in errors[:5]:  # Show max 5 errors
                messages.error(request, err)
            if len(errors) > 5:
                messages.error(request, f'... and {len(errors) - 5} more errors.')
        return redirect('dashboard_products')
    return redirect('dashboard_products')


@login_required(login_url='/dashboard/login/')
@user_passes_test(is_admin, login_url='/dashboard/login/')
def download_import_template(request):
    from .import_utils import generate_sample_excel
    from django.http import HttpResponse
    output = generate_sample_excel()
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="ishas_products_import_template.xlsx"'
    return response
