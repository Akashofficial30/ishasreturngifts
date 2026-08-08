from decimal import Decimal, InvalidOperation

import openpyxl
from products.models import Product, Category
from django.utils.text import slugify

# A guard against a spreadsheet with a stray million-row selection exhausting
# memory on a small instance.
MAX_IMPORT_ROWS = 5000


def cell_text(row_data, *keys, default=''):
    """Text from the first key present with a non-empty value.

    openpyxl yields None for an empty cell, and the key is still present, so
    dict.get(key, default) returns None rather than the default — str(None)
    then silently produced products named "None".
    """
    for key in keys:
        value = row_data.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return default


def parse_money(raw):
    """Decimal from a spreadsheet cell, tolerating '₹' and thousands commas."""
    if raw is None or not str(raw).strip():
        return None
    cleaned = str(raw).replace('₹', '').replace(',', '').strip()
    try:
        value = Decimal(cleaned)
    except (InvalidOperation, ValueError):
        return None
    # Decimal, not float: DecimalField columns should never inherit binary
    # floating point error from the import path.
    return value.quantize(Decimal('0.01')) if value >= 0 else None


def import_products_from_excel(file):
    """
    Import products from Excel file.
    Expected columns: Name, Category, Description, Price, Offer Price, Stock, Featured, Active
    Returns: (success_count, error_list)
    """
    errors = []
    success = 0

    try:
        # data_only=True returns the cached result of a formula rather than the
        # formula text, so a Price of "=B2*0.8" imports as a number.
        wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
        ws = wb.active
    except Exception as e:
        return 0, [f"Could not read Excel file: {str(e)}"]

    # Get headers from row 1
    headers = []
    for cell in ws[1]:
        headers.append(str(cell.value).strip().lower() if cell.value else '')

    required = ['name', 'category', 'price']
    for req in required:
        if req not in headers:
            return 0, [f"Missing required column: '{req}'. Required columns: Name, Category, Price"]

    # Process each row starting from row 2
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row_num - 1 > MAX_IMPORT_ROWS:
            errors.append(
                f"Stopped at row {row_num}: more than {MAX_IMPORT_ROWS} rows. "
                'Please split the file.'
            )
            break
        try:
            # `any(row)` treats a row of zeros as empty; test for actual content.
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue

            row_data = {}
            for i, val in enumerate(row):
                if i < len(headers) and headers[i]:
                    row_data[headers[i]] = val

            name = cell_text(row_data, 'name')
            if not name:
                errors.append(f"Row {row_num}: Name is empty — skipped")
                continue

            # Get or create category
            cat_name = cell_text(row_data, 'category', default='General')
            cat_slug = slugify(cat_name)
            if not cat_slug:
                # slugify strips non-Latin scripts entirely, and a blank slug
                # would collide every such category onto one row.
                errors.append(
                    f"Row {row_num}: Category '{cat_name}' has no usable name — skipped"
                )
                continue
            category, _ = Category.objects.get_or_create(
                slug=cat_slug,
                defaults={'name': cat_name, 'description': ''}
            )

            # Price
            price = parse_money(row_data.get('price'))
            if price is None:
                errors.append(f"Row {row_num}: Invalid price for '{name}' — skipped")
                continue

            # Offer price
            offer_price = parse_money(
                row_data.get('offer price')
                or row_data.get('offer_price')
                or row_data.get('offerprice')
            )
            if offer_price is not None and offer_price >= price:
                errors.append(
                    f"Row {row_num}: Offer price for '{name}' is not below the "
                    'price — ignored'
                )
                offer_price = None

            # Stock
            stock = 0
            stock_val = row_data.get('stock') or row_data.get('stock quantity') or row_data.get('quantity')
            if stock_val:
                try:
                    stock = max(0, int(float(str(stock_val).strip())))
                except (ValueError, TypeError):
                    stock = 0

            # Featured
            is_featured = cell_text(row_data, 'featured', default='no').lower() in (
                'yes', 'true', '1', 'y'
            )

            # Active
            is_active = cell_text(row_data, 'active', default='yes').lower() not in (
                'no', 'false', '0', 'n'
            )

            # Description
            description = cell_text(
                row_data, 'description', default=f'{name} - premium return gift'
            )

            # Generate unique slug. slugify can return '' for a name written in
            # a non-Latin script, and a product with an empty slug has no
            # reachable detail page.
            base_slug = slugify(name) or f'product-{row_num}'
            slug = base_slug
            counter = 1
            while Product.objects.filter(slug=slug).exists():
                slug = f"{base_slug}-{counter}"
                counter += 1

            # Create product
            Product.objects.create(
                name=name,
                slug=slug,
                category=category,
                description=description,
                price=price,
                offer_price=offer_price,
                stock_quantity=stock,
                is_featured=is_featured,
                is_active=is_active,
            )
            success += 1

        except Exception as e:
            errors.append(f"Row {row_num}: Error — {str(e)}")

    return success, errors


def generate_sample_excel():
    """Generate a sample Excel template for product import"""
    import io
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products Import"

    # Style
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="6B1F2A")
    center = Alignment(horizontal='center', vertical='center')
    thin = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    headers = ['Name', 'Category', 'Description', 'Price', 'Offer Price', 'Stock', 'Featured', 'Active']
    widths = [30, 20, 45, 10, 12, 8, 10, 8]

    for i, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 24

    # Sample rows
    samples = [
        ['Brass Deepam Set', 'Pooja Items', 'Traditional brass deepam set for pooja ceremonies. Ideal wedding return gift.', 450, 349, 200, 'Yes', 'Yes'],
        ['Sandalwood Soap Box', 'Wedding Return Gifts', 'Premium sandalwood soap set in a beautiful gift box.', 299, 249, 500, 'Yes', 'Yes'],
        ['Copper Water Bottle', 'Kitchen Essentials', 'Pure copper water bottle with Ayurvedic benefits.', 599, 499, 150, 'No', 'Yes'],
        ['Miniature Ganesha Idol', 'Home Décor', 'Beautifully crafted white marble finish Ganesha idol.', 350, '', 300, 'Yes', 'Yes'],
        ['Steel Tumbler Set', 'Kitchen Essentials', 'High quality stainless steel tumbler set of 2 pieces.', 350, 299, 350, 'No', 'Yes'],
    ]

    gold_fill = PatternFill("solid", fgColor="FFF8E1")
    for row_num, sample in enumerate(samples, 2):
        fill = gold_fill if row_num % 2 == 0 else None
        for col_num, val in enumerate(sample, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.border = thin
            cell.alignment = Alignment(vertical='center')
            if fill:
                cell.fill = fill
        ws.row_dimensions[row_num].height = 16

    # Instructions sheet
    ws2 = wb.create_sheet("Instructions")
    instructions = [
        ["IMPORT INSTRUCTIONS — Isha's Return Gifts"],
        [""],
        ["REQUIRED COLUMNS:"],
        ["  • Name — Product name (required)"],
        ["  • Category — Category name (will be created if doesn't exist)"],
        ["  • Price — Regular price in ₹ (required)"],
        [""],
        ["OPTIONAL COLUMNS:"],
        ["  • Description — Product description"],
        ["  • Offer Price — Discounted price (leave blank for no discount)"],
        ["  • Stock — Stock quantity (default: 0)"],
        ["  • Featured — Yes/No (shows on homepage)"],
        ["  • Active — Yes/No (visible on website, default: Yes)"],
        [""],
        ["TIPS:"],
        ["  • Keep headers exactly as shown in the Products sheet"],
        ["  • Do NOT delete the header row"],
        ["  • Leave Offer Price blank if no discount"],
        ["  • Images must be uploaded separately from admin"],
    ]
    ws2.column_dimensions['A'].width = 60
    for row_num, instruction in enumerate(instructions, 1):
        cell = ws2.cell(row=row_num, column=1, value=instruction[0] if instruction else '')
        if row_num == 1:
            cell.font = Font(bold=True, size=13, color="6B1F2A")
        elif instruction and instruction[0].endswith(':'):
            cell.font = Font(bold=True, size=11)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
